"""Typed run tables shared by CSV and openpyxl export, with explicit sheet maps."""
from dataclasses import dataclass, field
import csv
import json
import math
import re
from pathlib import Path
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_ROWS = 1_048_576
EXCEL_COLUMNS = 16_384
BODY_FONT = Font(name="Arial", size=10)
TEXT_ALIGNMENT = Alignment(vertical="center", horizontal="left")
NUMBER_ALIGNMENT = Alignment(vertical="center", horizontal="right")


def scalar(value):
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, np.ndarray):
        value = value.tolist()
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, allow_nan=False)
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("nonfinite table values must be explicitly marked unavailable with status")
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError("table text exceeds Excel's 32767-character cell limit; split the source field explicitly")
    return value


@dataclass
class RunTable:
    source_csv: str
    sheet_name: str
    columns: list[str]
    rows: list[tuple]
    units: dict[str, str] = field(default_factory=dict)
    kind: str = "empirical"

    @classmethod
    def from_records(cls, source_csv, sheet_name, records, *, units=None, kind="empirical"):
        records = list(records)
        columns = list(dict.fromkeys(k for row in records for k in row))
        return cls(source_csv, sheet_name, columns,
                   [tuple(str(row[k]) if k.endswith("_id") and row.get(k) is not None else scalar(row.get(k))
                          for k in columns) for row in records], units or {}, kind)

    @property
    def headers(self):
        return [f"{key} ({self.units[key]})" if key in self.units else key for key in self.columns]

    def write_csv(self, output):
        path = Path(output) / self.source_csv
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            if self.columns:
                writer.writerow(self.headers)
            writer.writerows(self.rows)


def configuration_table(config):
    records = []
    for key, value in config.items():
        if isinstance(value, (list, tuple)):
            records.extend(dict(parameter=f"{key}[{k}]", value=x, status="configured") for k, x in enumerate(value))
        else:
            records.append(dict(parameter=key, value=value, status="not_set" if value is None else "configured"))
    return RunTable.from_records("configuration.csv", "Configuration", records, kind="prescribed")


def _sheet_name(proposed, used):
    base = re.sub(r"[\\/*?:\[\]]", "_", proposed).strip("'") or "Table"
    candidate = base[:31]
    index = 2
    while candidate.casefold() in used:
        suffix = f"_{index}"
        candidate = base[:31-len(suffix)] + suffix
        index += 1
    used.add(candidate.casefold())
    return candidate


def _write_cell(sheet, row, column, value):
    value = scalar(value)
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"  # literal identifiers and '=...' text, never formulas
        cell.number_format = "@"
    elif isinstance(value, float):
        cell.number_format = "0.0000000000E+00" if value != 0 and abs(value) < 1e-5 else "0.0000000000"
    elif isinstance(value, int) and not isinstance(value, bool):
        if abs(value) > 999_999_999_999_999:
            raise ValueError("large integer exceeds Excel's exact 15-digit numeric precision; use a text identifier")
        cell.number_format = "0"
    cell.font = BODY_FONT
    cell.alignment = TEXT_ALIGNMENT if isinstance(value, str) else NUMBER_ALIGNMENT
    return cell


def _style_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="243C54")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 44
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 100) + 1)]
        width = min(48, max(14, max(map(len, values), default=0) + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width
        # Long provenance/status fields remain readable without rounding values.
        for row in range(2, min(sheet.max_row, 100) + 1):
            cell = sheet.cell(row, column)
            if isinstance(cell.value, str) and len(cell.value) > width - 2:
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
                height = min(409, 15 * math.ceil(len(cell.value) / max(10, width - 2)))
                sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 15, height)


def write_workbook(tables, output, *, filename="validation_results.xlsx", row_limit=EXCEL_ROWS,
                   column_limit=EXCEL_COLUMNS):
    """Split large tables in both dimensions, repeat headers, and record every part.

    Custom limits allow small tests of Excel-limit handling. Table values are
    snapshots of numerical outputs; Excel does not recompute simulation physics.
    """
    if not 2 <= row_limit <= EXCEL_ROWS or not 1 <= column_limit <= EXCEL_COLUMNS:
        raise ValueError("invalid workbook row/column limits")
    tables = list(tables)
    if len({t.source_csv for t in tables}) != len(tables):
        raise ValueError("each CSV source must be registered exactly once")
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    used = {"readme"}
    mapping = []
    for table in tables:
        if any(len(row) != len(table.columns) for row in table.rows):
            raise ValueError(f"inconsistent row length in {table.source_csv}")
        for column_start in range(0, max(1, len(table.columns)), column_limit):
            column_end = min(len(table.columns), column_start + column_limit)
            for row_start in range(0, max(1, len(table.rows)), row_limit - 1):
                row_end = min(len(table.rows), row_start + row_limit - 1)
                name = _sheet_name(table.sheet_name, used)
                sheet = workbook.create_sheet(name)
                headers = table.headers[column_start:column_end]
                for c, value in enumerate(headers, 1):
                    _write_cell(sheet, 1, c, value)
                for r, values in enumerate(table.rows[row_start:row_end], 2):
                    for c, value in enumerate(values[column_start:column_end], 1):
                        _write_cell(sheet, r, c, value)
                if not headers:
                    _write_cell(sheet, 1, 1, "Empty source table; no produced records")
                _style_sheet(sheet)
                mapping.append(dict(source_csv=table.source_csv, worksheet=name,
                                    first_data_row=row_start + 1, last_data_row=row_end,
                                    first_column=column_start + 1, last_column=column_end,
                                    source_rows=len(table.rows), source_columns=len(table.columns),
                                    quantity_kind=table.kind,
                                    units=json.dumps(table.units) if table.units else "See source headers / amplitude_unit / configuration",
                                    status="empty" if not table.rows else "exported"))
    notes = [
        ("Validation results", "CSV and Excel use the same typed table objects and row ordering."),
        ("Numerical source", "Full waveform arrays are in waveform.npz; only requested CSV tables enter Excel."),
        ("Null convention", "Blank cells represent unavailable values. Accompanying status fields explain missing measurements."),
        ("Precision", "Numeric cells retain floating-point values; formatting controls display. NPZ is the complete numerical source."),
        ("Identifiers", "Identifiers and bit strings are literal text. Text starting with '=' is never a formula."),
        ("Units", "Unit-bearing headers or amplitude_unit columns identify amplitude units. Variances use amplitude units squared; time uses s or UI as labeled."),
        ("Prescribed", "Inputs configured by the simulation, including nominal levels and per-level endpoint amplitude sigmas."),
        ("Theoretical", "Analytic independent-endpoint references, before impairments. Adjacent segments share endpoints and are correlated."),
        ("Empirical", "Statistics use actual symbol/transition labels. One observation per segment per phase; classes are not pooled."),
        ("Fitted", "Only included when produced. No new GMM fitter or eye-boundary measurement is fabricated."),
        ("Original lane sigma", "Independent-lane benchmark parameter only; it is not interpreted as two endpoint sigmas."),
        ("Shared mode", "Amplitude continuity is guaranteed; derivative continuity is not guaranteed. Endpoint fluctuations are amplitude variations, not timing jitter."),
    ]
    # All rows remain a flat, filterable metadata table with the CSV mapping.
    readme_rows = [dict(section="Conventions", source_csv=key, worksheet=value) for key, value in notes]
    readme_rows.extend(dict(section="CSV mapping", **entry) for entry in mapping)
    keys = list(dict.fromkeys(k for row in readme_rows for k in row))
    for c, key in enumerate(keys, 1):
        _write_cell(readme, 1, c, key)
    for r, row in enumerate(readme_rows, 2):
        for c, key in enumerate(keys, 1):
            _write_cell(readme, r, c, row.get(key))
    _style_sheet(readme)
    readme.column_dimensions["C"].width = 115
    for r in range(2, len(notes) + 2):
        readme.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="center")
        readme.row_dimensions[r].height = 34
    path = Path(output) / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    (path.parent / "workbook_mapping.json").write_text(json.dumps(mapping, indent=2), encoding="utf-8")
    return mapping


def verify_workbook(tables, output, mapping):
    """Reopen and compare all exported table cells, including types and nulls."""
    workbook = load_workbook(Path(output) / "validation_results.xlsx", data_only=False, read_only=True)
    by_source = {t.source_csv: t for t in tables}
    count = 0
    for entry in mapping:
        table = by_source[entry["source_csv"]]
        sheet = workbook[entry["worksheet"]]
        first, last = entry["first_column"] - 1, entry["last_column"]
        if not table.columns:
            continue
        expected = [table.headers[first:last]] + [row[first:last] for row in table.rows[entry["first_data_row"]-1:entry["last_data_row"]]]
        if sheet.max_row != len(expected) or sheet.max_column != last - first:
            raise AssertionError(f"truncated or extra cells in {sheet.title}")
        for row_number, (actual, values) in enumerate(zip(sheet.iter_rows(), expected), 1):
            for cell, value in zip(actual, values):
                value = scalar(value)
                if isinstance(value, float):
                    if not isinstance(cell.value, (int, float)) or not math.isclose(cell.value, value, rel_tol=2e-15, abs_tol=1e-300):
                        raise AssertionError(f"numeric mismatch at {sheet.title}!{cell.coordinate}")
                elif value == "":
                    if cell.value not in (None, ""):
                        raise AssertionError("empty-text mismatch")
                elif cell.value != value:
                    raise AssertionError(f"value mismatch at {sheet.title}!{cell.coordinate}")
                if isinstance(value, str) and value and cell.data_type != "s":
                    raise AssertionError("text identifier or literal text lost its type")
                count += 1
    workbook.close()
    return count
