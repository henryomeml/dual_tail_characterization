from dataclasses import replace
import csv
import json
from pathlib import Path
import numpy as np
import pytest
from openpyxl import load_workbook
from dual_tail_characterization import Config, generate_waveform, extract_eye, paired_eyes
from dual_tail_characterization.shared_endpoints import (
    effective_lanes, logistic_weight, normalized_duration, kappa_for_duration, lane_statistics,
)
from dual_tail_characterization.lane_benchmark import original_lane_records
from dual_tail_characterization.shared_endpoint_demo import build_validation_tables, run_demo
from dual_tail_characterization.validation_workbook import RunTable, write_workbook, verify_workbook


def shared(**kwargs):
    return Config(waveform_mode="shared_endpoints", **kwargs)


def test_original_table_is_exact_and_complete():
    original = original_lane_records()
    with (Path(__file__).parents[1] / "benchmarks/customized_pam4_independent_lanes.csv").open() as file:
        saved = [{k: float(v) for k, v in row.items()} for row in csv.DictReader(file)]
    assert saved == original
    assert len(original) == 16
    assert {(r["start"], r["end"]) for r in original} == {(a,b) for a in range(4) for b in range(4)}
    assert original[11] == dict(start=2, end=3, mean_start=2, mean_end=2.97, sigma=.1)
    assert original[10]["mean_end"] == 2.97  # preserve unusual same-level benchmark row


def test_effective_definitions_and_arbitrary_pam_n():
    levels, sigmas = [-.2, .75, 2.9], [.3, 0, .1]
    lanes = effective_lanes(levels, sigmas)
    assert len(lanes) == 9
    for row in lanes:
        assert row["mean_start"] == levels[row["start"]]
        assert row["mean_end"] == levels[row["end"]]
        assert row["sigma_start"] == sigmas[row["start"]]
        assert row["sigma_end"] == sigmas[row["end"]]
    w = generate_waveform(shared(levels=levels, endpoint_sigmas=sigmas, n_symbols=30))
    assert np.all(np.isfinite(w.final))


def test_reproducibility_and_stream_separation():
    c = shared(n_symbols=100, endpoint_sigmas=[.2, .1, .1, .2])
    a, b = generate_waveform(c), generate_waveform(c)
    for name in ("symbol_indices", "symbol_endpoint_values", "endpoint_standard_normals", "final"):
        np.testing.assert_array_equal(getattr(a, name), getattr(b, name))
    deterministic = generate_waveform(replace(c, endpoint_sigmas=[0]*4))
    np.testing.assert_array_equal(a.symbol_indices, deterministic.symbol_indices)
    np.testing.assert_array_equal(a.endpoint_standard_normals, deterministic.endpoint_standard_normals)
    np.testing.assert_array_equal(a.symbol_endpoint_values, a.symbol_amplitudes + np.asarray(c.endpoint_sigmas)[a.symbol_indices]*a.endpoint_standard_normals)


@pytest.mark.parametrize("samples_per_ui", [7, 128])
@pytest.mark.parametrize("phase", [0, .123, -.123])
def test_exact_endpoint_reuse_clock_context_and_extraction(samples_per_ui, phase):
    c = shared(n_symbols=18, samples_per_ui=samples_per_ui, endpoint_sigmas=[.2, .1, .1, .2])
    w = generate_waveform(c)
    assert len(w.final) == 17*samples_per_ui + 1
    assert np.all(np.diff(w.time) > 0)
    np.testing.assert_array_equal(w.final[::samples_per_ui], w.symbol_endpoint_values)
    np.testing.assert_allclose(w.time[::samples_per_ui], w.symbol_centers, rtol=1e-15)
    np.testing.assert_allclose(w.symbol_boundaries/c.symbol_period, np.arange(19))
    one, two = paired_eyes(w, phase)
    np.testing.assert_array_equal(one.traces, two.traces[:, :samples_per_ui+1])
    np.testing.assert_allclose(one.start_times/c.symbol_period, one.clock_start_symbols-.5+phase)
    for first, trace in zip(two.source_start_indices, two.traces):
        np.testing.assert_array_equal(trace, np.interp(first+np.arange(len(trace)), np.arange(len(w.final)), w.final))
        assert first >= 0 and first + 2*samples_per_ui <= len(w.final)-1
    if phase == 0:
        np.testing.assert_array_equal(one.context_symbol_indices[0], w.symbol_indices[:2])
        np.testing.assert_array_equal(two.context_symbol_indices[0], w.symbol_indices[:3])
        np.testing.assert_array_equal(two.traces[:, samples_per_ui], w.symbol_endpoint_values[two.clock_start_symbols])
        adjacent = extract_eye(w, starts=two.clock_start_symbols+1)
        np.testing.assert_array_equal(one.traces[:, -1], adjacent.traces[:, 0])
        assert one.start_times[0] == w.symbol_centers[0]
    with pytest.raises(ValueError, match="valid interval"):
        extract_eye(w, starts=[0])


@pytest.mark.parametrize("duration", [.05, .2, .6, .79])
def test_normalized_duration_is_solved(duration):
    kappa = kappa_for_duration(duration)
    assert normalized_duration(kappa) == pytest.approx(duration, abs=1e-12)
    u = np.array([.5-duration/2, .5, .5+duration/2])
    np.testing.assert_allclose(logistic_weight(u, kappa), [.1, .5, .9], atol=1e-12)
    np.testing.assert_array_equal(logistic_weight([0, 1], kappa), [0, 1])
    assert not np.isclose(kappa_for_duration(.6), 2*np.log(9)/.6, rtol=1e-4)


def test_small_kappa_and_zero_sigma_curves():
    np.testing.assert_allclose(logistic_weight([0, .1, .5, 1], 1e-300), [0, .1, .5, 1])
    assert normalized_duration(1e-300) == .8
    c = shared(n_symbols=3, symbol_mode="explicit", sequence=[0, 3, 3], endpoint_sigmas=[0]*4)
    w = generate_waveform(c)
    weight = logistic_weight(np.arange(129)/128, w.metadata["effective_kappa"])
    np.testing.assert_allclose(w.final[:129], 3*weight, atol=1e-15)
    np.testing.assert_array_equal(w.final[128:], np.full(129, 3.))


def test_lane_statistics_and_same_symbol_ensemble():
    mean, variance, sigma = lane_statistics(0, 3, .2, .1, .5, 22.)
    assert mean == 1.5
    assert variance == (.2**2+.1**2)/4
    assert sigma == np.sqrt(variance)
    c = shared(n_symbols=30_000, symbol_mode="explicit", sequence=[1]*30_000, endpoint_sigmas=[.2, .1, .1, .2], samples_per_ui=8)
    w = generate_waveform(c)
    midpoint = w.final[4::8]
    assert midpoint.mean() == pytest.approx(1, abs=.003)
    assert midpoint.var() == pytest.approx(.1**2/2, rel=.04)
    assert w.symbol_endpoint_values.var() == pytest.approx(.1**2, rel=.04)
    assert np.any(np.diff(w.final[:9]) != 0)  # same-symbol realizations are not forced flat
    assert np.corrcoef(midpoint[:-1], midpoint[1:])[0,1] == pytest.approx(.5, abs=.04)


def test_all_empirical_lanes_and_endpoints_match_theory():
    w = generate_waveform(shared(n_symbols=32_000, samples_per_ui=8, endpoint_sigmas=[.2, .1, .1, .2]))
    tables, _, checks = build_validation_tables(w)
    assert all(row["status"] in ("pass", "not_implemented") for row in checks)
    lane_table = next(t for t in tables if t.source_csv == "lane_empirical.csv")
    assert len(lane_table.rows) == 16*3
    counts = next(t for t in tables if t.source_csv == "transition_counts.csv")
    assert all(row[3] > 1500 for row in counts.rows)


@pytest.mark.parametrize("kwargs", [dict(endpoint_sigmas=[.1]), dict(endpoint_sigmas=[-.1]*4),
    dict(shared_kappa=0), dict(shared_kappa=float("inf")), dict(channel_bandwidth_hz=1e9),
    dict(noise_std=.01), dict(n_symbols=1), dict(transition_10_90_ui=.9)])
def test_invalid_shared_settings(kwargs):
    with pytest.raises(ValueError):
        shared(**kwargs)


def test_workbook_limits_types_nulls_and_all_cells(tmp_path):
    table = RunTable("typed.csv", "illegal/[name]"*4,
        ["identifier", "bits", "amplitude", "missing_numeric", "status"],
        [(f"=literal_{i}", "0010", .1234567890123456+i, None, "unavailable") for i in range(7)],
        units={"amplitude": "V"})
    table.write_csv(tmp_path)
    mapping = write_workbook([table], tmp_path, row_limit=4, column_limit=2)
    assert len(mapping) == 9
    assert verify_workbook([table], tmp_path, mapping) > 0
    book = load_workbook(tmp_path / "validation_results.xlsx")
    assert len(set(s.casefold() for s in book.sheetnames)) == len(book.sheetnames)
    for item in mapping:
        sheet = book[item["worksheet"]]
        assert len(sheet.title) <= 31
        assert sheet.max_row <= 4 and sheet.max_column <= 2
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
    first = book[mapping[0]["worksheet"]]
    assert first["A2"].value == "=literal_0" and first["A2"].data_type == "s"
    assert first["B2"].value == "0010"
    book.close()


def test_demo_outputs_and_every_csv_is_in_workbook(tmp_path):
    run_demo(tmp_path, n_symbols=2000)
    mapping = json.loads((tmp_path / "workbook_mapping.json").read_text())
    assert {m["source_csv"] for m in mapping} == {p.name for p in tmp_path.glob("*.csv")}
    for filename in ("waveform.npz", "waveform.png", "eye_1ui.png", "eye_2ui.png", "eye_1ui_traces.png", "eye_2ui_traces.png",
                     "endpoint_level_histograms.png", "transition_conditioned_histograms.png", "validation_results.xlsx", "REPORT.md"):
        assert (tmp_path / filename).stat().st_size > 0
    with np.load(tmp_path / "waveform.npz", allow_pickle=False) as archive:
        np.testing.assert_array_equal(archive["symbol_endpoint_values"], archive["final"][::128])
        assert "symbol_centers" in archive
    with np.load(tmp_path / "validation_histogram_counts.npz") as hist:
        assert hist["endpoint_counts"].sum() == 2000
        assert hist["transition_phase_counts"].sum() == 1999*3
    workbook = load_workbook(tmp_path / "validation_results.xlsx", data_only=False)
    for item in mapping:
        sheet = workbook[item["worksheet"]]
        with (tmp_path / item["source_csv"]).open(newline="", encoding="utf-8") as file:
            csv_rows = list(csv.reader(file))
        assert len(csv_rows) == sheet.max_row
        assert csv_rows[0] == [cell.value for cell in sheet[1]]
        for csv_row, cells in zip(csv_rows[1:], sheet.iter_rows(min_row=2)):
            for text, cell in zip(csv_row, cells):
                if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                    assert float(text) == pytest.approx(cell.value, rel=2e-15, abs=1e-300)
                elif cell.value is None:
                    assert text == ""
                elif isinstance(cell.value, bool):
                    assert text == str(cell.value)
                else:
                    assert text == cell.value
    workbook.close()
